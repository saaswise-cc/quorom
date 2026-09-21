"""The configuration the rest of the suite cannot reach: exactly one CRM.

The suite runs with both CRMs unconfigured (conftest enforces it), and a real
deployment runs with one. Every "compares two CRMs" defect so far has lived in
that gap — a column or flag that exists to set one CRM against the other, left
rendering when there is no other. Fixed one display at a time, it recurred.

So this drives the whole weekly sequence after the database — reconcile,
coverage, the stakeholder list, the workbook and the HTML view — with stubbed
adapters in each one-CRM configuration, and asserts that nothing comparing two
CRMs appears on any tab or on the page. The both-configured case is here too, so
the gate is shown to open as well as close.

No real Config() anywhere: it reads the environment, and a failing assertion
can put a live credential in a traceback. The config is a stub carrying only
the fields these steps read.
"""

from __future__ import annotations

import datetime as dt
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from quorom.crm.contact import Contact
from quorom.weekly import coverage as coverage_mod
from quorom.weekly import people as people_mod
from quorom.weekly import stakeholders as stakeholders_mod
from quorom.weekly import view as view_mod
from quorom.weekly import workbook as workbook_mod

# Everything in the output that exists to compare one CRM with the other.
# Each must appear only when both are configured.
CROSS_CRM_HEADERS = ("In HubSpot?", "In Salesforce?")
CROSS_CRM_FLAGS = ("title only in", "title differs")

PROFILE = {
    "employee_count_min": 50,
    "employee_count_max": 5000,
    "hq_geographies": ["North America"],
    "focus_seniority": ["vp"],
}

TODAY = dt.date.today().isoformat()


class _SF:
    """Salesforce as the weekly steps call it. Unconfigured, it answers the
    way the real adapter does: no record, no stats, no firmographics, no
    bench."""

    linkedin_available = True

    def __init__(self, configured: bool) -> None:
        self.configured = configured

    def contact_by_email(self, email):
        if not self.configured:
            return None
        return {
            # A title in this CRM only — the case "title only in" fired on.
            "dana@acme.example": Contact(
                name="Dana Reyes", title="VP Sales", email=email,
                linkedin="https://www.linkedin.com/in/dana-reyes",
            ),
            # A title in both CRMs, differing — the case "title differs" fires on.
            "lee@acme.example": Contact(name="Lee Park", title="VP Marketing", email=email),
        }.get(email)

    def domain_stats(self, domain, terms):
        if not self.configured:
            return {"sf_total": 0, "sf_senior": 0, "account_id": None}
        return {"sf_total": 3, "sf_senior": 1, "account_id": "001x"}

    def account_firmographics(self, account_id):
        out = {"name": "", "employees": "", "hq": "", "account_type": "",
               "country": "", "city": "", "state": ""}
        if not self.configured or not account_id:
            return out
        return {**out, "name": "Acme", "employees": 500,
                "hq": "Austin, TX, United States", "country": "United States"}

    def senior_bench(self, domain, terms):
        if not self.configured:
            return []
        return [
            Contact(name="Dana Reyes", title="VP Sales", email="dana@acme.example",
                    mobile=True, linkedin="https://www.linkedin.com/in/dana-reyes",
                    last_activity=TODAY),
        ]


class _HS:
    """HubSpot as the weekly steps call it, with the same unconfigured answers
    as the real adapter."""

    def __init__(self, configured: bool) -> None:
        self.configured = configured

    def contact_by_email(self, email):
        if not self.configured:
            return None
        return {
            # A title in this CRM only.
            "sam@acme.example": Contact(name="Sam Fox", title="Head of Ops", email=email),
            # Differs from the Salesforce title above.
            "lee@acme.example": Contact(name="Lee Park", title="CMO", email=email),
        }.get(email)

    def count_domain(self, domain):
        return 7 if self.configured else 0


def _cfg(sf_on: bool, hs_on: bool):
    return SimpleNamespace(
        salesforce=SimpleNamespace(configured=sf_on),
        hubspot=SimpleNamespace(configured=hs_on),
        customer_account_types=(),
        recent_days=90,
        group_call_min=8,
        shortlist_size=3,
    )


def _people():
    rows = [
        {"attendee_name": "Dana Reyes", "email": "dana@acme.example",
         "domain": "acme.example", "meeting_title": "Intro"},
        {"attendee_name": "Sam Fox", "email": "sam@acme.example",
         "domain": "acme.example", "meeting_title": "Intro"},
        {"attendee_name": "Lee Park", "email": "lee@acme.example",
         "domain": "acme.example", "meeting_title": "Intro"},
        # In neither CRM, so it lands on tab 2 whichever one is configured.
        {"attendee_name": "Ari Stone", "email": "ari@acme.example",
         "domain": "acme.example", "meeting_title": "Intro"},
    ]
    return people_mod.dedupe_people(rows)


def _run(tmp_path, sf_on: bool, hs_on: bool):
    """The weekly sequence from step 3 on, as run.py orders it."""
    cfg = _cfg(sf_on, hs_on)
    sf, hs = _SF(sf_on), _HS(hs_on)

    people = _people()
    companies = people_mod.group_companies(people)
    reconciled = [people_mod.reconcile(p, sf, hs) for p in people]
    coverage = coverage_mod.build_coverage(
        cfg, companies, PROFILE, sf, hs, log=lambda *_: None
    )
    terms = coverage_mod.seniority_terms(PROFILE)
    stakeholders, _ = stakeholders_mod.build(cfg, coverage, terms, {}, sf)

    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(
        cfg, reconciled, coverage, [], stakeholders, xlsx,
        profile=PROFILE, geo_label="North America",
    )
    html_path = view_mod.render(xlsx, "example.com")
    return load_workbook(xlsx), open(html_path).read()


def _headers(ws):
    return [h for h in next(ws.iter_rows(max_row=1, values_only=True)) if h]


def _cells(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    yield ws.title, str(value)


@pytest.mark.parametrize(
    "sf_on, hs_on, crm",
    [(True, False, "Salesforce"), (False, True, "HubSpot")],
    ids=["salesforce-only", "hubspot-only"],
)
def test_one_crm_renders_nothing_that_compares_two(tmp_path, sf_on, hs_on, crm):
    wb, html = _run(tmp_path, sf_on, hs_on)

    # Not vacuous: the configured CRM did answer with titles, which is what the
    # "title only in" flag used to fire on — on every such row.
    titles = [
        r[2] for r in wb["1 - Met this week"].iter_rows(min_row=2, values_only=True)
        if r[2]
    ]
    assert titles, "no CRM title rendered — the flag assertions below prove nothing"

    for ws in wb.worksheets:
        for header in CROSS_CRM_HEADERS:
            assert header not in _headers(ws), f"{header!r} on {ws.title}"
    for sheet, value in _cells(wb):
        for flag in CROSS_CRM_FLAGS:
            assert flag not in value, f"{flag!r} on {sheet}: {value!r}"
    for text in CROSS_CRM_HEADERS + CROSS_CRM_FLAGS:
        assert text not in html

    # Tab 2 has no Source column. The CRM it was checked against is stated once,
    # in the caption — on the sheet and on the page.
    ws2 = wb["2 - Missing from CRM"]
    assert _headers(ws2) == ["Name", "Email", "Company (domain)", "Flag"]
    missing = [r for r in ws2.iter_rows(min_row=2, values_only=True) if r[0] and any(r[1:])]
    assert missing, "nobody on tab 2 — the caption has no rows to explain"
    assert f"Checked against {crm}." in [c for _, c in _cells(wb)]
    assert f"Checked against {crm}." in html


def test_both_crms_still_compare(tmp_path):
    """The gate opens as well as closes: with both configured, the comparison
    columns and flags are exactly what they were."""
    wb, html = _run(tmp_path, True, True)

    ws1 = wb["1 - Met this week"]
    flag = _headers(ws1).index("Flag")
    flags = {r[0]: r[flag] for r in ws1.iter_rows(min_row=2, values_only=True) if r[0]}
    assert flags["Dana Reyes"] == "title only in Salesforce"
    assert flags["Sam Fox"] == "title only in HubSpot"
    assert flags["Lee Park"] == "title differs (SF: VP Marketing / HS: CMO)"

    ws2 = wb["2 - Missing from CRM"]
    assert _headers(ws2) == [
        "Name", "Email", "Company (domain)", "In HubSpot?", "In Salesforce?", "Flag",
    ]
    # The columns say it per row, so no caption repeats it.
    assert not any(v.startswith("Checked against") for _, v in _cells(wb))
    assert "Checked against" not in html
