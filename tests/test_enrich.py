"""The enrichment pass, end to end after the database, with a stubbed provider.

Covers the configured and unconfigured states from the first commit: every
earlier "not configured collapsed into false" defect came from a state with no
test. The provider here is invented — the pass reads only what the `enrich`
package promises, so any provider behaves the same way to it.

No real Config(): it reads the environment, and a failing assertion can put a
live credential in a traceback. Stubs throughout.
"""

from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from quorom import enrich
from quorom.crm.contact import Contact
from quorom.enrich import Company, Person
from quorom.weekly import coverage as coverage_mod
from quorom.weekly import enrichment
from quorom.weekly import people as people_mod
from quorom.weekly import stakeholders as stakeholders_mod
from quorom.weekly import view as view_mod
from quorom.weekly import workbook as workbook_mod

PROFILE = {
    "employee_count_min": 50,
    "employee_count_max": 5000,
    "hq_geographies": ["North America"],
    "focus_seniority": ["vp", "c-level"],
}

# What the CRM holds about each company. Globex is recorded at 10 employees —
# the misrecorded-headcount case that removes a company from the map unseen.
FIRMOGRAPHICS = {
    "acme.example": {"name": "Acme", "employees": 500, "country": "United States"},
    "globex.example": {"name": "Globex", "employees": 10, "country": "United States"},
    "initech.example": {"name": "Initech", "employees": 300, "country": "Canada"},
}

BENCH = {
    "acme.example": [
        Contact(name="Dana Reyes", title="VP Sales", email="dana@acme.example",
                linkedin="https://www.linkedin.com/in/dana-reyes"),
        Contact(name="Lee Park", title="VP Marketing", email="lee@acme.example"),
    ],
    # The CRM's LinkedIn URL for Kim points at someone else — the wrong-URL case.
    "globex.example": [Contact(name="Kim Lo", title="CRO", email="kim@globex.example",
                               linkedin="https://www.linkedin.com/in/kimlo")],
    # Ray's email finds nothing; his CRM LinkedIn URL finds him, somewhere else.
    "initech.example": [Contact(name="Ray Oh", title="VP Ops", email="ray@initech.example",
                                linkedin="https://linkedin.com/in/ray-oh/")],
}

IN_CRM = {"dana@acme.example", "lee@acme.example"}


class _SF:
    """Salesforce-only, the configuration the first deployment runs."""

    configured = True
    linkedin_available = True

    def contact_by_email(self, email):
        for bench in BENCH.values():
            for c in bench:
                if c.email == email and email in IN_CRM:
                    return c
        return None

    def domain_stats(self, domain, terms):
        return {"sf_total": 2, "sf_senior": 1, "account_id": domain}

    def account_firmographics(self, account_id):
        f = FIRMOGRAPHICS.get(account_id, {})
        return {"name": f.get("name", ""), "employees": f.get("employees", ""),
                "hq": f.get("country", ""), "country": f.get("country", ""),
                "account_type": "", "city": "", "state": ""}

    def senior_bench(self, domain, terms):
        return list(BENCH.get(domain, []))


class _HS:
    configured = False

    def contact_by_email(self, email):
        return None

    def count_domain(self, domain):
        return 0


class _Provider:
    """An invented provider with the interface the `enrich` package promises."""

    display_name = "Example"

    def __init__(self):
        self.person_calls: list[str] = []
        self.linkedin_calls: list[str] = []
        self.company_calls: list[str] = []
        self.people = {
            # Still at Acme, with a different title there — but an advisory
            # seat elsewhere is listed first. Still here, not moved.
            "dana@acme.example": Person(
                name="Dana Reyes", title="Advisor", employer_name="Board Co",
                employer_domain="board.example",
                linkedin="https://www.linkedin.com/in/dana-reyes",
                current_jobs=(("board.example", "Board Co", "Advisor"),
                              ("acme.example", "Acme", "SVP Sales")),
            ),
            # Moved on.
            "lee@acme.example": Person(name="Lee Park", title="CMO", employer_name="Globex",
                                       employer_domain="globex.example", updated="2026-06-01"),
            # Not in the CRM; the provider knows them.
            "ari@acme.example": Person(name="Ari Stone", title="Director of Ops",
                                       employer_name="Acme", employer_domain="acme.example",
                                       linkedin="https://www.linkedin.com/in/ari-stone"),
        }
        self.companies = {
            # Agrees with the CRM's verdict, though not its number.
            "acme.example": Company(name="Acme", domain="acme.example", employees=520,
                                    country="United States"),
            # Disagrees: the CRM says 10 employees, so the CRM says no.
            "globex.example": Company(name="Globex", domain="globex.example",
                                      employees=150, country="United States"),
        }

    def person_by_email(self, email):
        self.person_calls.append(email)
        return self.people.get(email)

    # The provider accepts on the handle alone; the name check is the pass's.
    by_linkedin = {
        "ray-oh": Person(name="Ray Oh", title="COO", employer_name="Hooli",
                         employer_domain="hooli.example",
                         linkedin="https://www.linkedin.com/in/ray-oh"),
        "kimlo": Person(name="Kimberly Stone", title="CRO", employer_name="Other Co",
                        employer_domain="other.example",
                        linkedin="https://www.linkedin.com/in/kimlo"),
    }

    def person_by_linkedin(self, url):
        from quorom.enrich import linkedin_handle

        self.linkedin_calls.append(url)
        return self.by_linkedin.get(linkedin_handle(url))

    def company_by_domain(self, domain):
        self.company_calls.append(domain)
        return self.companies.get(domain)


def _cfg():
    return SimpleNamespace(
        salesforce=SimpleNamespace(configured=True),
        hubspot=SimpleNamespace(configured=False),
        customer_account_types=(),
        recent_days=90,
        group_call_min=8,
        shortlist_size=3,
    )


def _attendees():
    rows = [
        {"attendee_name": "Dana Reyes", "email": "dana@acme.example", "domain": "acme.example"},
        {"attendee_name": "Lee Park", "email": "lee@acme.example", "domain": "acme.example"},
        {"attendee_name": "", "email": "ari@acme.example", "domain": "acme.example"},
        {"attendee_name": "Support", "email": "support@acme.example", "domain": "acme.example"},
        {"attendee_name": "Kim Lo", "email": "kim@globex.example", "domain": "globex.example"},
        {"attendee_name": "Pat Vo", "email": "pat@initech.example", "domain": "initech.example"},
    ]
    return people_mod.dedupe_people([{**r, "meeting_title": "Intro"} for r in rows])


def _run(tmp_path, provider):
    """Steps 3 to 6 in run.py's order, with the enrichment pass where run.py
    puts it: companies before the map is chosen, people after."""
    cfg, sf, hs = _cfg(), _SF(), _HS()
    people = _attendees()
    reconciled = [people_mod.reconcile(p, sf, hs) for p in people]
    coverage = coverage_mod.build_coverage(
        cfg, people_mod.group_companies(people), PROFILE, sf, hs, log=lambda *_: None
    )
    pass_ = enrichment.start(provider)
    queue = []
    if pass_:
        enrichment.companies(pass_, coverage, PROFILE)
    rows, _ = stakeholders_mod.build(
        cfg, coverage, coverage_mod.seniority_terms(PROFILE), {}, sf
    )
    if pass_:
        enrichment.stakeholders(pass_, rows)
        enrichment.not_in_crm(pass_, reconciled)
        queue = enrichment.review_queue(pass_, coverage, rows)

    xlsx = str(tmp_path / "weekly_stakeholder_map_2026-08-17.xlsx")
    workbook_mod.build_workbook(
        cfg, reconciled, coverage, [], rows, xlsx, profile=PROFILE,
        geo_label="North America",
        enrichment=provider.display_name if provider else None, queue=queue,
    )
    html = open(view_mod.render(xlsx, "example.com")).read()
    return load_workbook(xlsx), html, coverage


def _table(ws):
    """Header row and data rows as dicts, captions and blank rows dropped."""
    rows = list(ws.iter_rows(values_only=True))
    headers = [h for h in rows[0] if h]
    out = []
    for r in rows[1:]:
        # A caption sits alone in column A; a data row has something after it.
        # A blank Name is a real row — the case the provider's name fills.
        if any(v not in (None, "") for v in r[1:]):
            out.append(dict(zip(headers, r)))
    return headers, out


# --- off ---------------------------------------------------------------------- #


def test_with_no_provider_nothing_changes(tmp_path):
    wb, html, _ = _run(tmp_path, None)

    assert wb.sheetnames == [
        "1 - Met this week", "2 - Company coverage", "3 - Stakeholder list",
    ]
    for ws in wb.worksheets:
        headers, _ = _table(ws)
        assert not any("Example" in h or h in ("Profile check", "Still at company?")
                       for h in headers), ws.title
    # The CRM's "no" stands alone, so the misrecorded company is off the map —
    # the invisible error the pass exists to surface.
    _, rows = _table(wb["3 - Stakeholder list"])
    assert not any("Globex" in str(r["Company"]) for r in rows)
    assert "Review queue" not in html


# --- on ----------------------------------------------------------------------- #


def test_company_verdicts_are_compared_not_numbers(tmp_path):
    wb, _, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["2 - Company coverage"])
    by = {r["Company"]: r for r in rows}

    assert headers[-3:] == ["Employees (Example)", "HQ (Example)", "Profile check"]
    # 500 against 520 is inside the band either way: not a finding.
    assert by["acme.example"]["Profile check"] == "agrees"
    assert by["globex.example"]["Profile check"] == "disputed — Example says yes"
    # Looked, and the provider has nothing. Stated, never blank.
    assert by["initech.example"]["Profile check"] == "not found in Example"
    assert by["initech.example"]["Employees (Example)"] is None
    # The CRM's own value is never replaced.
    assert by["globex.example"]["Employees"] == 10


def test_a_disputed_rejection_reaches_the_map_marked(tmp_path):
    wb, _, _ = _run(tmp_path, _Provider())
    _, rows = _table(wb["3 - Stakeholder list"])

    globex = [r for r in rows if str(r["Company"]).startswith("Globex")]
    assert globex and all(r["Company"] == "Globex (profile disputed)" for r in globex)
    # Agreeing companies are not marked.
    assert any(r["Company"] == "Acme" for r in rows)


def test_still_at_company_and_differences(tmp_path):
    wb, _, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["3 - Stakeholder list"])
    by = {r["Name"]: r for r in rows}

    assert headers[-3:] == ["Still at company?", "Title (Example)", "LinkedIn (Example)"]
    # Acme is among Dana's current positions, though not listed first.
    assert by["Dana Reyes"]["Still at company?"] == "yes"
    assert by["Lee Park"]["Still at company?"] == "no — now at Globex"
    # Kim's CRM LinkedIn URL points at someone else: not used, still not found.
    assert by["Kim Lo"]["Still at company?"] == "not found in Example"
    # Ray's email found nothing; his LinkedIn URL did, and the name agrees.
    assert by["Ray Oh"]["Still at company?"] == "no — now at Hooli (matched on LinkedIn)"
    # Shown only where it differs — and it is the title at *this* company, not
    # the first one listed. The CRM's title stays in Title.
    assert by["Dana Reyes"]["Title (Example)"] == "SVP Sales"
    assert by["Dana Reyes"]["Title"] == "VP Sales"
    assert by["Dana Reyes"]["LinkedIn (Example)"] in (None, "")
    # A detected move flags the row; it is not removed.
    assert "Lee Park" in by
    # Their provider title is for the new job, so it is not set beside this one.
    assert by["Lee Park"]["Title (Example)"] in (None, "")


def test_people_not_in_the_crm_get_a_name_title_and_linkedin(tmp_path):
    """On Met this week, for the people the CRM does not hold — including the
    LinkedIn URL, which is what someone needs to connect with them."""
    wb, html, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["1 - Met this week"])
    by = {r["Email"]: r for r in rows}

    # The provider's name is in the header; the caption says what that means.
    assert "The Example columns are enrichment from Example, not your CRM" in html

    assert ["Name (Example)", "Title (Example)", "LinkedIn (Example)"] == [
        h for h in headers if "(Example)" in h
    ]
    assert by["ari@acme.example"]["Name (Example)"] == "Ari Stone"
    assert by["ari@acme.example"]["Title (Example)"] == "Director of Ops"
    assert by["ari@acme.example"]["LinkedIn (Example)"] == "https://www.linkedin.com/in/ari-stone"
    assert by["pat@initech.example"]["Name (Example)"] == "not found in Example"
    assert by["support@acme.example"]["Name (Example)"] == "not looked up — shared inbox"
    # People the CRM does hold are not given provider columns here: the
    # stakeholder list is where the provider is set beside a CRM record.
    assert by["dana@acme.example"]["Name (Example)"] in (None, "")


def test_each_person_and_company_is_looked_up_once_and_inboxes_never(tmp_path):
    provider = _Provider()
    _run(tmp_path, provider)

    assert "support@acme.example" not in provider.person_calls
    assert len(provider.person_calls) == len(set(provider.person_calls))
    # LinkedIn is tried only where the email found nothing and the CRM holds a
    # URL: Kim and Ray here — never Dana or Lee, whose emails matched.
    assert sorted(provider.linkedin_calls) == [
        "https://linkedin.com/in/ray-oh/", "https://www.linkedin.com/in/kimlo",
    ]
    assert len(provider.company_calls) == len(set(provider.company_calls))


def test_the_review_queue_holds_what_a_person_should_settle(tmp_path):
    wb, html, _ = _run(tmp_path, _Provider())
    headers, rows = _table(wb["4 - Review queue"])

    assert headers == ["What", "Company", "Person", "CRM says", "Example says", "Check"]
    kinds = [r["What"] for r in rows]
    # Most consequential first: a disputed verdict can hide a whole company.
    assert kinds[0] == "Profile fit disputed"
    assert "May have left" in kinds
    assert "Title differs" in kinds
    assert "Headcount or HQ missing" in kinds          # the provider had nothing
    moved = next(r for r in rows if r["What"] == "May have left" and r["Person"] == "Lee Park")
    assert moved["Example says"].startswith("now at Globex")
    ray = next(r for r in rows if r["What"] == "May have left" and r["Person"] == "Ray Oh")
    assert "(matched on LinkedIn)" in ray["Example says"]
    # A handle match under another name becomes a question, not a finding.
    wrong = next(r for r in rows if r["What"] == "CRM LinkedIn may be someone else")
    assert wrong["Person"] == "Kim Lo" and "Kimberly Stone" in wrong["Example says"]
    # The move is the finding; a title at the new company is not a disagreement.
    assert not any(r["What"] == "Title differs" and r["Person"] == "Lee Park" for r in rows)
    # Every row says where to check.
    assert all(r["Check"] for r in rows)

    assert "Review queue" in html
    assert '<span class="fl">disputed — Example says yes</span>' in html
    assert '<span class="fl">no — now at Globex</span>' in html


def test_one_crm_with_a_provider_still_compares_no_two_crms(tmp_path):
    """The one-CRM defect class, with a provider added: nothing comparing two
    CRMs appears when only one is configured."""
    wb, _, _ = _run(tmp_path, _Provider())
    for ws in wb.worksheets:
        headers, _ = _table(ws)
        assert "In HubSpot?" not in headers and "In Salesforce?" not in headers
        for row in ws.iter_rows(values_only=True):
            for v in row:
                assert "title only in" not in str(v or "")


def test_two_configured_providers_are_refused(monkeypatch):
    class _A:
        display_name, env_vars = "A", ("A_KEY",)

        @classmethod
        def from_env(cls, environ=None):
            return cls()

    class _B(_A):
        display_name, env_vars = "B", ("B_KEY",)

    monkeypatch.setattr(enrich, "_modules", lambda: iter([_A, _B]))
    with pytest.raises(enrich.MoreThanOneProvider):
        enrich.configured({})


# --- run.py wiring, against the real database -------------------------------- #


def test_the_weekly_run_checks_the_provider_first_and_adds_tab_5(
    database, gong_calls, tmp_path, monkeypatch
):
    """The order is the point: the provider's free check runs before the
    database is read or a CRM called, and the company pass before the map is
    chosen. No CRM is configured here (conftest), so every verdict is the
    CRM's "not assessed" beside the provider's.

    The config is a stub, not Config(): it borrows Config's two week methods
    and nothing else, so nothing is read from the environment.
    """
    from quorom.config import Config
    from quorom.weekly import run as run_mod
    from tests.test_import_and_weekly import (
        ACCOUNT, _import, _seed_account, _seed_profile,
    )

    account_id = _seed_account(database)
    _import(database, account_id, gong_calls)
    _seed_profile(database, account_id)

    provider = _Provider()
    provider.companies["acme.com"] = Company(name="Acme", domain="acme.com",
                                             employees=900, country="United States")
    events: list[str] = []
    provider.check = lambda: events.append("check") or "stub plan, 100 credits available"
    monkeypatch.setattr(run_mod.enrich, "configured", lambda: provider)
    real_connect = run_mod.db.connect
    monkeypatch.setattr(run_mod.db, "connect",
                        lambda cfg: (events.append("db"), real_connect(cfg))[1])

    cfg = SimpleNamespace(
        database_url=database, account=ACCOUNT, week_start="2026-08-17",
        output_dir=str(tmp_path), tz_offset="-04", shortlist_size=3, group_call_min=8,
        recent_days=90, customer_account_types=(), retain_runs=False,
        gong=SimpleNamespace(configured=False),
        salesforce=SimpleNamespace(
            configured=False, uses_client_credentials=False, access_token="",
            instance_url="", token_url="", client_id="", client_secret="",
            api_version="v61.0",
        ),
        hubspot=SimpleNamespace(configured=False, api_key=""),
    )
    cfg.week_bounds = lambda: Config.week_bounds(cfg)
    cfg.week_days_remaining = lambda: Config.week_days_remaining(cfg)
    logged: list[str] = []
    paths = run_mod.run_weekly(cfg, log=logged.append)

    assert events[:2] == ["check", "db"]
    assert any("Enrichment: Example" in line for line in logged)

    wb = load_workbook(paths["xlsx"])
    assert "4 - Review queue" in wb.sheetnames
    _, rows = _table(wb["2 - Company coverage"])
    acme = next(r for r in rows if r["Company"] == "acme.com")
    assert acme["Profile check"] == "CRM not assessed — Example says yes"
    # Not assessed is not disputed: there is no CRM verdict to disagree with.
    assert "(profile disputed)" not in json.dumps(
        [c.value for row in wb["3 - Stakeholder list"].iter_rows() for c in row], default=str
    )

    dump = json.loads(open(paths["json"]).read())
    assert dump["enrichment_provider"] == "Example"
    assert isinstance(dump["review_queue"], list)


def test_a_provider_without_a_linkedin_lookup_is_not_asked(tmp_path):
    """The LinkedIn lookup is optional in the provider interface. Without it,
    an email miss stays a miss, and nothing else changes."""

    class _EmailOnly:
        display_name = "Example"

        def __init__(self):
            self._inner = _Provider()

        def person_by_email(self, email):
            return self._inner.person_by_email(email)

        def company_by_domain(self, domain):
            return self._inner.company_by_domain(domain)

    wb, _, _ = _run(tmp_path, _EmailOnly())
    _, rows = _table(wb["3 - Stakeholder list"])
    by = {r["Name"]: r for r in rows}
    assert by["Ray Oh"]["Still at company?"] == "not found in Example"
    _, queue = _table(wb["4 - Review queue"])
    assert not any(r["What"] == "CRM LinkedIn may be someone else" for r in queue)


@pytest.mark.parametrize(
    "a, b, agree",
    [
        ("Dana Reyes", "Dana Reyes", True),
        ("Dana M. Reyes", "dana reyes", True),       # middle initial, case
        ("José Álvarez", "Jose Alvarez", True),       # accents
        ("Mary-Kate O'Neil", "Mary Kate ONeil", True),   # hyphen, apostrophe
        ("Dana Reyes", "Dana Rivera", False),
        ("Dana Reyes", "Kimberly Stone", False),
        ("", "Dana Reyes", False),
    ],
)
def test_names_agree_on_first_and_last(a, b, agree):
    assert enrichment.names_agree(a, b) is agree


def test_the_stakeholder_list_says_how_it_is_built(tmp_path):
    """Which companies, which people, how many and in what order — in the
    reader's own values. A list of names with no stated rule reads as a
    recommendation from nowhere."""
    wb, html, _ = _run(tmp_path, _Provider())
    captions = [
        r[0] for r in wb["3 - Stakeholder list"].iter_rows(min_row=2, values_only=True)
        if r[0] and not any(r[1:])
    ]
    rule = captions[0]
    assert "50–5,000 employees, HQ in North America" in rule
    assert "plus any whose profile fit is disputed" in rule   # Globex is
    assert "up to 3 people from your CRM at VP or C-suite level" in rule
    assert "most senior first, then most recently contacted" in rule
    assert "People not in your CRM cannot appear here" in rule
    assert rule in html
